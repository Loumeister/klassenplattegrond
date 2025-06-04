import csv
import random
import argparse # Added for CLI
import sys # For exit()
from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont

# Main script for the Seating Chart Generator

def parse_comma_separated_string(name_string):
    """
    Parses a comma-separated string of names into a list.

    Args:
        name_string: A string containing names separated by commas.
                     Example: "Alice, Bob, Charlie "

    Returns:
        A list of strings, where each string is a cleaned name.
        Returns an empty list if the input string is empty or only whitespace.
    """
    if not name_string or name_string.isspace():
        return []
    names = [name.strip() for name in name_string.split(',')]
    return [name for name in names if name] # Remove any empty strings that might result from "Alice,,Bob"

def parse_csv_file(file_path):
    """
    Parses a CSV file to extract names from the first column.

    Args:
        file_path: The path to the CSV file.

    Returns:
        A list of names extracted from the first column of the CSV.
        Returns an empty list if the file is not found or if an error occurs,
        and prints an error message to stderr.
    """
    names = []
    try:
        with open(file_path, mode='r', newline='', encoding='utf-8') as csvfile:
            reader = csv.reader(csvfile)
            for row in reader:
                if row and row[0].strip(): # Check if row is not empty and first element is not empty
                    names.append(row[0].strip())
    except FileNotFoundError:
        print(f"Error: File not found at {file_path}", file=sys.stderr)
        return []
    except Exception as e:
        print(f"An error occurred while parsing CSV file {file_path}: {e}", file=sys.stderr)
        return []
    return names

def parse_tsv_file(file_path):
    """
    Parses a TSV file to extract names from the first column.

    Args:
        file_path: The path to the TSV file.

    Returns:
        A list of names extracted from the first column of the TSV.
        Returns an empty list if the file is not found or if an error occurs,
        and prints an error message to stderr.
    """
    names = []
    try:
        with open(file_path, mode='r', newline='', encoding='utf-8') as tsvfile:
            reader = csv.reader(tsvfile, delimiter='\t')
            for row in reader:
                if row and row[0].strip(): # Check if row is not empty and first element is not empty
                    names.append(row[0].strip())
    except FileNotFoundError:
        print(f"Error: File not found at {file_path}", file=sys.stderr)
        return []
    except Exception as e:
        print(f"An error occurred while parsing TSV file {file_path}: {e}", file=sys.stderr)
        return []
    return names

def shuffle_names(names_list):
    """
    Shuffles a list of names.

    Args:
        names_list: A list of names.

    Returns:
        A new list with the same names in a random order.
        The original list is not modified.
    """
    shuffled_list = names_list[:] # Create a copy
    random.shuffle(shuffled_list)
    return shuffled_list

def arrange_seats(student_names, teacher_name=None):
    """
    Arranges student names into a predefined seating structure.

    The structure is 3 tables, each with 5 columns, and each column has 2 seats.
    Total student seats: 3 * 5 * 2 = 30.

    Args:
        student_names: A list of student names (should be pre-shuffled).
        teacher_name: An optional string for the teacher's name.

    Returns:
        A dictionary with two keys:
        'teacher': The teacher's name (or None).
        'tables': A list of 3 tables. Each table is a list of 5 columns.
                  Each column is a list of 2 student names (or None if seat is empty).
    """
    seating_chart = {
        'teacher': teacher_name,
        'tables': []
    }

    num_tables = 3
    num_columns_per_table = 5
    num_seats_per_column = 2
    total_student_seats = num_tables * num_columns_per_table * num_seats_per_column

    seated_students = student_names[:total_student_seats]
    if len(seated_students) < total_student_seats:
        seated_students.extend([None] * (total_student_seats - len(seated_students)))

    student_idx = 0
    for _ in range(num_tables):
        table_data = []
        for _ in range(num_columns_per_table):
            column_data = []
            for _ in range(num_seats_per_column):
                if student_idx < len(seated_students):
                    column_data.append(seated_students[student_idx])
                    student_idx += 1
                else:
                    column_data.append(None)
            table_data.append(column_data)
        seating_chart['tables'].append(table_data)
    return seating_chart

def generate_excel_output(seating_data, output_filename):
    """
    Generates an Excel file with two sheets: Teacher's View and Student's View.

    Args:
        seating_data: The dictionary from `arrange_seats`.
        output_filename: The name for the Excel file (e.g., "seating_chart.xlsx").
    """
    wb = Workbook()
    thin_border_side = Side(border_style="thin", color="000000")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    ws_teacher = wb.active
    ws_teacher.title = "Teacher's View"

    teacher_name = seating_data.get('teacher')
    if teacher_name:
        teacher_cell = ws_teacher['A1']
        teacher_cell.value = f"Teacher: {teacher_name}"
        teacher_cell.font = Font(bold=True, size=14)
        ws_teacher.merge_cells('A1:E1')
        teacher_cell.alignment = Alignment(horizontal='center')

    current_row_teacher = 3

    for table_idx, table_data in enumerate(seating_data['tables']):
        ws_teacher.cell(row=current_row_teacher, column=1, value=f"Table {table_idx + 1}").font = Font(bold=True)
        current_row_teacher += 1

        for col_num_idx in range(len(table_data)):
            ws_teacher.cell(row=current_row_teacher, column=col_num_idx + 1, value=f"Col {col_num_idx+1}").font = Font(italic=True)
        current_row_teacher +=1

        for seat_row_idx in range(2):
            for col_idx, column_data in enumerate(table_data):
                cell = ws_teacher.cell(row=current_row_teacher + seat_row_idx, column=col_idx + 1)
                cell.value = column_data[seat_row_idx]
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                ws_teacher.column_dimensions[get_column_letter(col_idx + 1)].width = 15
                ws_teacher.row_dimensions[current_row_teacher + seat_row_idx].height = 30
        current_row_teacher += 2
        current_row_teacher += 1

    ws_student = wb.create_sheet("Student's View")
    if teacher_name:
        teacher_row_student_view = 16
        student_teacher_cell = ws_student.cell(row=teacher_row_student_view, column=1)
        student_teacher_cell.value = f"Teacher: {teacher_name}"
        student_teacher_cell.font = Font(bold=True, size=14)
        ws_student.merge_cells(start_row=teacher_row_student_view, start_column=1, end_row=teacher_row_student_view, end_column=5)
        student_teacher_cell.alignment = Alignment(horizontal='center')

    current_row_student = 1
    reversed_tables = seating_data['tables'][::-1]

    for table_idx, table_data in enumerate(reversed_tables):
        ws_student.cell(row=current_row_student, column=1, value=f"Your Table Area {table_idx + 1} (Front is Table 3)").font = Font(bold=True)
        current_row_student += 1

        for col_num_idx in range(len(table_data)):
            ws_student.cell(row=current_row_student, column=len(table_data) - col_num_idx, value=f"Col {col_num_idx+1}").font = Font(italic=True)
        current_row_student +=1

        for seat_row_idx in range(2):
            actual_seat_row_idx = 1 - seat_row_idx
            for col_idx, column_data in enumerate(table_data):
                actual_col_idx = len(table_data) - 1 - col_idx
                original_student_name = reversed_tables[table_idx][actual_col_idx][actual_seat_row_idx]
                cell = ws_student.cell(row=current_row_student + seat_row_idx, column=col_idx + 1)
                cell.value = original_student_name
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                ws_student.column_dimensions[get_column_letter(col_idx + 1)].width = 15
                ws_student.row_dimensions[current_row_student + seat_row_idx].height = 30
        current_row_student += 2
        current_row_student += 1

    try:
        wb.save(output_filename)
        # Message printed in main block
    except Exception as e:
        print(f"Error saving Excel file {output_filename}: {e}", file=sys.stderr)

def generate_jpg_output(seating_data, output_filename):
    """
    Generates a JPG image with Teacher's View and Student's View (rotated).
    """
    CELL_WIDTH, CELL_HEIGHT = 100, 40
    TABLE_MARGIN, VIEW_SPACING, PAGE_MARGIN = 20, 40, 20
    TEXT_PADDING = 5
    LINE_COLOR, TEXT_COLOR, BACKGROUND_COLOR = "black", "black", "white"
    FONT_SIZE_NAMES, FONT_SIZE_LABELS, FONT_SIZE_TEACHER = 14, 16, 18
    NUM_TABLES, NUM_COLS_PER_TABLE, NUM_ROWS_PER_TABLE = 3, 5, 2

    try:
        font_names = ImageFont.truetype("arial.ttf", FONT_SIZE_NAMES)
        font_labels = ImageFont.truetype("arial.ttf", FONT_SIZE_LABELS)
        font_teacher = ImageFont.truetype("arial.ttf", FONT_SIZE_TEACHER)
    except IOError:
        print("Arial font not found, using default font.", file=sys.stderr)
        font_names = ImageFont.load_default()
        font_labels = ImageFont.load_default()
        font_teacher = ImageFont.load_default()

    teacher_name_height = FONT_SIZE_TEACHER + TEXT_PADDING * 2 if seating_data.get('teacher') else 0
    table_label_height = FONT_SIZE_LABELS + TEXT_PADDING
    one_view_content_width = NUM_COLS_PER_TABLE * CELL_WIDTH
    one_table_visual_height = table_label_height + NUM_ROWS_PER_TABLE * CELL_HEIGHT
    all_tables_height = NUM_TABLES * one_table_visual_height + (NUM_TABLES - 1) * TABLE_MARGIN
    one_view_content_height = teacher_name_height + all_tables_height + TEXT_PADDING
    one_view_height = one_view_content_height + 2 * PAGE_MARGIN
    img_width = one_view_content_width + 2 * PAGE_MARGIN
    img_height = (2 * one_view_height) + VIEW_SPACING

    img = Image.new('RGB', (img_width, img_height), BACKGROUND_COLOR)
    draw = ImageDraw.Draw(img)

    def draw_single_view(offset_x_base, offset_y_base, is_student_view):
        teacher_name = seating_data.get('teacher')
        tables_data = seating_data['tables']
        current_y = offset_y_base + PAGE_MARGIN

        if teacher_name:
            name_w, name_h = draw.textbbox((0,0), f"Teacher: {teacher_name}", font=font_teacher)[2:4]
            if not is_student_view:
                teacher_x = offset_x_base + (one_view_content_width - name_w) / 2 + PAGE_MARGIN
                draw.text((teacher_x, current_y), f"Teacher: {teacher_name}", fill=TEXT_COLOR, font=font_teacher)
                current_y += name_h + TEXT_PADDING * 2

        processed_tables = tables_data[::-1] if is_student_view else tables_data
        for table_idx, table_content in enumerate(processed_tables):
            table_display_idx = NUM_TABLES - table_idx if is_student_view else table_idx + 1
            label_text = f"Table {table_display_idx}" if not is_student_view else f"Your Table Area {table_idx + 1} (Front is Original Table 3)"
            label_w, label_h = draw.textbbox((0,0), label_text, font=font_labels)[2:4]
            label_x = offset_x_base + (one_view_content_width - label_w) / 2 + PAGE_MARGIN
            draw.text((label_x, current_y), label_text, fill=TEXT_COLOR, font=font_labels)
            current_y += label_h + TEXT_PADDING
            table_start_x = offset_x_base + PAGE_MARGIN
            for col_idx_disp in range(NUM_COLS_PER_TABLE): # Displayed column index
                for row_idx_disp in range(NUM_ROWS_PER_TABLE): # Displayed row index
                    data_col_idx = NUM_COLS_PER_TABLE - 1 - col_idx_disp if is_student_view else col_idx_disp
                    data_row_idx = NUM_ROWS_PER_TABLE - 1 - row_idx_disp if is_student_view else row_idx_disp
                    student_name = table_content[data_col_idx][data_row_idx]
                    cell_x1 = table_start_x + col_idx_disp * CELL_WIDTH
                    cell_y1 = current_y + row_idx_disp * CELL_HEIGHT
                    cell_x2 = cell_x1 + CELL_WIDTH
                    cell_y2 = cell_y1 + CELL_HEIGHT
                    draw.rectangle([cell_x1, cell_y1, cell_x2, cell_y2], outline=LINE_COLOR)
                    if student_name:
                        text_w, text_h = draw.textbbox((0,0), student_name, font=font_names)[2:4]
                        text_x = cell_x1 + (CELL_WIDTH - text_w) / 2
                        text_y = cell_y1 + (CELL_HEIGHT - text_h) / 2
                        draw.text((text_x, text_y), student_name, fill=TEXT_COLOR, font=font_names)
            current_y += NUM_ROWS_PER_TABLE * CELL_HEIGHT + TABLE_MARGIN

        if teacher_name and is_student_view:
            current_y -= TABLE_MARGIN # Adjust for last table margin
            current_y += TEXT_PADDING * 2
            name_w, name_h = draw.textbbox((0,0), f"Teacher: {teacher_name}", font=font_teacher)[2:4]
            teacher_x = offset_x_base + (one_view_content_width - name_w) / 2 + PAGE_MARGIN
            draw.text((teacher_x, current_y), f"Teacher: {teacher_name}", fill=TEXT_COLOR, font=font_teacher)

    draw_single_view(0, 0, is_student_view=False)
    student_view_offset_y = one_view_height + VIEW_SPACING
    draw_single_view(0, student_view_offset_y, is_student_view=True)

    try:
        img.save(output_filename, "JPEG")
        # Message printed in main block
    except Exception as e:
        print(f"Error saving JPG file {output_filename}: {e}", file=sys.stderr)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generate a seating chart from a list of names.")

    parser.add_argument(
        "input_source",
        help="Path to the input file (CSV/TSV) or a comma-separated string of names."
    )
    parser.add_argument(
        "--input-type", "-t",
        required=True,
        choices=['list', 'csv', 'tsv'],
        help="Type of the input source."
    )
    parser.add_argument(
        "--output-format", "-f",
        required=True,
        choices=['excel', 'jpg'],
        help="Desired output format."
    )
    parser.add_argument(
        "--output-file", "-o",
        required=True,
        help="Name and path for the output file (e.g., chart.xlsx or chart.jpg)."
    )
    parser.add_argument(
        "--teacher",
        default=None,
        help="Optional name of the teacher."
    )

    args = parser.parse_args()

    student_names = []
    if args.input_type == 'list':
        student_names = parse_comma_separated_string(args.input_source)
    elif args.input_type == 'csv':
        student_names = parse_csv_file(args.input_source)
    elif args.input_type == 'tsv':
        student_names = parse_tsv_file(args.input_source)
    else:
        # Should not happen due to argparse choices
        print(f"Error: Invalid input type '{args.input_type}'.", file=sys.stderr)
        sys.exit(1)

    if not student_names:
        print("Error: No names found from input or input was empty. Exiting.", file=sys.stderr)
        sys.exit(1)

    print(f"Successfully parsed {len(student_names)} names.")

    shuffled_student_names = shuffle_names(student_names)
    print("Names shuffled.")

    seating_arrangement = arrange_seats(shuffled_student_names, args.teacher)
    print("Seating arrangement created.")

    if args.output_format == 'excel':
        generate_excel_output(seating_arrangement, args.output_file)
    elif args.output_format == 'jpg':
        generate_jpg_output(seating_arrangement, args.output_file)
    else:
        # Should not happen due to argparse choices
        print(f"Error: Invalid output format '{args.output_format}'.", file=sys.stderr)
        sys.exit(1)

    print(f"Successfully generated {args.output_file}")
